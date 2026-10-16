"""A reader's own space: folders of rulings with notes, alerts, history, downloads.

Folders belong to a person; a folder marked shared is seen by everyone in the same
organization -- a firm's shared research. Alerts are saved searches: the worker emails
what is new since the last one. History is read from client_activity, which every view
and search already writes.

Downloads are counted, limited per day by the plan, and marked: each Word file served
carries, at the foot of every page, who downloaded it, when, and a reference that
leads back to the row here. A file passed around is then a file that can be traced.
Reading at a pace no person reads at -- hundreds of rulings an hour -- is refused.
"""
from __future__ import annotations

import datetime as dt
import io
import json
import time
import uuid
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel
from sqlalchemy import Date, DateTime, Integer, String, Text, UniqueConstraint, func, or_, select
from sqlalchemy.orm import Mapped, mapped_column

import db
from deps import STAFF, current_user


class Folder(db.Base):
    __tablename__ = "folders"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    owner: Mapped[str] = mapped_column(String, index=True)
    org_id: Mapped[Optional[int]] = mapped_column(Integer, nullable=True, index=True)
    name: Mapped[str] = mapped_column(String, default="")
    shared: Mapped[int] = mapped_column(Integer, default=0)
    created_at: Mapped[dt.datetime] = mapped_column(DateTime, default=dt.datetime.utcnow)


class FolderItem(db.Base):
    __tablename__ = "folder_items"
    __table_args__ = (UniqueConstraint("folder_id", "doc_id", name="uq_folder_item"),)
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    folder_id: Mapped[int] = mapped_column(Integer, index=True)
    doc_id: Mapped[str] = mapped_column(String, index=True)
    note: Mapped[str] = mapped_column(Text, default="")
    added_by: Mapped[str] = mapped_column(String, default="")
    added_at: Mapped[dt.datetime] = mapped_column(DateTime, default=dt.datetime.utcnow)


class Alert(db.Base):
    __tablename__ = "alerts"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    user_email: Mapped[str] = mapped_column(String, index=True)
    name: Mapped[str] = mapped_column(String, default="")
    query: Mapped[str] = mapped_column(String, default="")
    filters: Mapped[str] = mapped_column(String, default="{}")   # chamber, theme, year, city
    lang: Mapped[str] = mapped_column(String, default="ar")
    frequency: Mapped[str] = mapped_column(String, default="weekly")   # daily | weekly
    active: Mapped[int] = mapped_column(Integer, default=1)
    last_sent_at: Mapped[Optional[dt.datetime]] = mapped_column(DateTime, nullable=True)
    created_at: Mapped[dt.datetime] = mapped_column(DateTime, default=dt.datetime.utcnow)


class ReadingTime(db.Base):
    """Seconds a reader kept a ruling open and in view, per day, per language. A page
    sends a beat every half minute while it is visible; a click that leaves at once
    sends nothing, so an opened ruling and a read one can be told apart."""
    __tablename__ = "reading_time"
    __table_args__ = (UniqueConstraint("actor", "doc_id", "day", "lang", name="uq_reading_time"),)
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    actor: Mapped[str] = mapped_column(String, index=True)
    doc_id: Mapped[str] = mapped_column(String, index=True)
    day: Mapped[dt.date] = mapped_column(Date, index=True)
    lang: Mapped[str] = mapped_column(String, default="ar")
    seconds: Mapped[int] = mapped_column(Integer, default=0)


class Download(db.Base):
    __tablename__ = "downloads"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    ref: Mapped[str] = mapped_column(String, unique=True)
    user_email: Mapped[str] = mapped_column(String, index=True)
    org_id: Mapped[Optional[int]] = mapped_column(Integer, nullable=True, index=True)
    doc_id: Mapped[str] = mapped_column(String)
    at: Mapped[dt.datetime] = mapped_column(DateTime, default=dt.datetime.utcnow, index=True)


# ---------------- pace and quota ----------------
VIEWS_PER_HOUR = 300
_views: dict = {}


def note_view(user: dict) -> None:
    """Refuse a customer's account reading faster than a person reads."""
    if user["role"] in STAFF:
        return
    now = time.time()
    hits = [t for t in _views.get(user["email"], []) if now - t < 3600]
    if len(hits) >= VIEWS_PER_HOUR:
        db.add_audit(user["email"], "throttled", "user", user["email"], new=f"{len(hits)} views/h")
        raise HTTPException(status_code=429, detail={"pace": VIEWS_PER_HOUR})
    hits.append(now)
    _views[user["email"]] = hits


def downloads_today(email: str) -> int:
    start = dt.datetime.combine(dt.date.today(), dt.time.min)
    with db.session() as s:
        return s.scalar(select(func.count()).select_from(Download).where(
            Download.user_email == email, Download.at >= start)) or 0


def start_download(user: dict, doc_id: str) -> str:
    """Check the day's quota and record the download. Returns its reference."""
    if user["role"] not in STAFF:
        import billing
        quota = billing.daily_downloads(user.get("org_id"))
        used = downloads_today(user["email"])
        if used >= quota:
            raise HTTPException(status_code=429, detail={"quota": quota, "used": used})
    ref = uuid.uuid4().hex[:10].upper()
    with db.session() as s:
        s.add(Download(ref=ref, user_email=user["email"], org_id=user.get("org_id"), doc_id=doc_id))
        s.commit()
    return ref


def mark_docx(data: bytes, line: str) -> bytes:
    """The same Word file with `line` at the foot of every page. A file that will not
    open as Word comes back unchanged rather than failing the download."""
    try:
        from docx import Document
        from docx.shared import Pt
        d = Document(io.BytesIO(data))
        for section in d.sections:
            p = section.footer.add_paragraph()
            r = p.add_run(line)
            r.font.size = Pt(7)
        out = io.BytesIO()
        d.save(out)
        return out.getvalue()
    except Exception:
        return data


# ---------------- routes ----------------
router = APIRouter()


class FolderBody(BaseModel):
    name: str
    shared: bool = False


class ItemBody(BaseModel):
    doc_id: str
    note: str = ""


class NoteBody(BaseModel):
    note: str


def _visible_folders(user: dict):
    cond = Folder.owner == user["email"]
    if user.get("org_id"):
        cond = or_(cond, (Folder.org_id == user["org_id"]) & (Folder.shared == 1))
    return cond


def _folder(s, folder_id: int, user: dict, write: bool = False) -> Folder:
    f = s.get(Folder, folder_id)
    if not f:
        raise HTTPException(status_code=404, detail="Not found")
    mine = f.owner == user["email"]
    shared = f.shared and f.org_id and f.org_id == user.get("org_id")
    if not (mine or shared):
        raise HTTPException(status_code=404, detail="Not found")
    if write and not (mine or shared):
        raise HTTPException(status_code=403, detail="Not allowed")
    return f


def _folder_dict(f: Folder, count: int, user: dict) -> dict:
    return {"id": f.id, "name": f.name, "shared": bool(f.shared), "owner": f.owner,
            "mine": f.owner == user["email"], "count": count,
            "created_at": f.created_at.isoformat()}


@router.get("/api/me/folders")
def folders(user: dict = Depends(current_user)) -> list:
    with db.session() as s:
        rows = list(s.scalars(select(Folder).where(_visible_folders(user)).order_by(Folder.name)))
        counts = dict(s.execute(select(FolderItem.folder_id, func.count())
                                .where(FolderItem.folder_id.in_([f.id for f in rows]))
                                .group_by(FolderItem.folder_id)).all()) if rows else {}
        return [_folder_dict(f, counts.get(f.id, 0), user) for f in rows]


@router.post("/api/me/folders")
def folder_create(body: FolderBody, user: dict = Depends(current_user)) -> dict:
    if not body.name.strip():
        raise HTTPException(status_code=400, detail="Name required")
    with db.session() as s:
        f = Folder(owner=user["email"], org_id=user.get("org_id"), name=body.name.strip()[:120],
                   shared=int(bool(body.shared and user.get("org_id"))))
        s.add(f)
        s.commit()
        return _folder_dict(f, 0, user)


@router.patch("/api/me/folders/{folder_id}")
def folder_update(folder_id: int, body: FolderBody, user: dict = Depends(current_user)) -> dict:
    with db.session() as s:
        f = _folder(s, folder_id, user)
        if f.owner != user["email"]:
            raise HTTPException(status_code=403, detail="Only the owner renames a folder")
        f.name = body.name.strip()[:120] or f.name
        f.shared = int(bool(body.shared and f.org_id))
        s.commit()
        n = s.scalar(select(func.count()).select_from(FolderItem).where(FolderItem.folder_id == f.id)) or 0
        return _folder_dict(f, n, user)


@router.delete("/api/me/folders/{folder_id}")
def folder_delete(folder_id: int, user: dict = Depends(current_user)) -> dict:
    with db.session() as s:
        f = _folder(s, folder_id, user)
        if f.owner != user["email"]:
            raise HTTPException(status_code=403, detail="Only the owner deletes a folder")
        for it in s.scalars(select(FolderItem).where(FolderItem.folder_id == folder_id)):
            s.delete(it)
        s.delete(f)
        s.commit()
    return {"ok": True}


@router.get("/api/me/folders/{folder_id}")
def folder_detail(folder_id: int, lang: str = "ar", user: dict = Depends(current_user)) -> dict:
    with db.session() as s:
        f = _folder(s, folder_id, user)
        items = list(s.scalars(select(FolderItem).where(FolderItem.folder_id == folder_id)
                               .order_by(FolderItem.added_at.desc())))
        out = _folder_dict(f, len(items), user)
        ids = [i.doc_id for i in items]
        decs = {d.doc_id: d for d in s.scalars(select(db.Decision).where(db.Decision.doc_id.in_(ids)))} if ids else {}
        extra = db.i18n_for(ids, lang) if lang != "ar" else {}
        out["items"] = []
        for i in items:
            d = decs.get(i.doc_id)
            if not d or d.state != "published":
                continue
            out["items"].append({
                "id": i.id, "doc_id": i.doc_id, "note": i.note, "added_by": i.added_by,
                "added_at": i.added_at.isoformat(), "decision_no": d.decision_no,
                "date": d.decision_date, "chamber": d.category, "theme": d.theme or "",
                "abstract": (d.rule if len(d.rule or "") > 80 else d.abstract) or "",
                "abstract_fr": (extra.get(i.doc_id) or {}).get("abstract", "")})
        return out


@router.post("/api/me/folders/{folder_id}/items")
def item_add(folder_id: int, body: ItemBody, user: dict = Depends(current_user)) -> dict:
    d = db.get_decision(body.doc_id)
    if not d or d["state"] != "published":
        raise HTTPException(status_code=404, detail="Ruling not found")
    with db.session() as s:
        _folder(s, folder_id, user, write=True)
        if s.scalar(select(FolderItem).where(FolderItem.folder_id == folder_id,
                                              FolderItem.doc_id == body.doc_id)):
            return {"ok": True, "already": True}
        s.add(FolderItem(folder_id=folder_id, doc_id=body.doc_id, note=body.note[:4000],
                         added_by=user["email"]))
        s.commit()
    return {"ok": True}


@router.patch("/api/me/folders/{folder_id}/items/{item_id}")
def item_note(folder_id: int, item_id: int, body: NoteBody, user: dict = Depends(current_user)) -> dict:
    with db.session() as s:
        _folder(s, folder_id, user, write=True)
        it = s.get(FolderItem, item_id)
        if not it or it.folder_id != folder_id:
            raise HTTPException(status_code=404, detail="Not found")
        it.note = body.note[:4000]
        s.commit()
    return {"ok": True}


@router.delete("/api/me/folders/{folder_id}/items/{item_id}")
def item_remove(folder_id: int, item_id: int, user: dict = Depends(current_user)) -> dict:
    with db.session() as s:
        _folder(s, folder_id, user, write=True)
        it = s.get(FolderItem, item_id)
        if it and it.folder_id == folder_id:
            s.delete(it)
            s.commit()
    return {"ok": True}


@router.get("/api/me/rulings/{doc_id}/folders")
def ruling_folders(doc_id: str, user: dict = Depends(current_user)) -> list:
    """Which of my folders already hold this ruling."""
    with db.session() as s:
        ids = [f.id for f in s.scalars(select(Folder).where(_visible_folders(user)))]
        if not ids:
            return []
        return list(s.scalars(select(FolderItem.folder_id).where(
            FolderItem.folder_id.in_(ids), FolderItem.doc_id == doc_id)))


# -- alerts --
class AlertBody(BaseModel):
    name: str = ""
    query: str = ""
    chamber: str = ""
    theme: str = ""
    year: str = ""
    city: str = ""
    lang: str = "ar"
    frequency: str = "weekly"
    active: Optional[bool] = None


def _alert(a: Alert) -> dict:
    f = json.loads(a.filters or "{}")
    return {"id": a.id, "name": a.name, "query": a.query, "lang": a.lang,
            "frequency": a.frequency, "active": bool(a.active), **{k: f.get(k, "") for k in
                                                                    ("chamber", "theme", "year", "city")},
            "last_sent_at": a.last_sent_at.isoformat() if a.last_sent_at else "",
            "created_at": a.created_at.isoformat()}


@router.get("/api/me/alerts")
def alerts(user: dict = Depends(current_user)) -> list:
    with db.session() as s:
        return [_alert(a) for a in s.scalars(select(Alert).where(Alert.user_email == user["email"])
                                             .order_by(Alert.created_at.desc()))]


@router.post("/api/me/alerts")
def alert_create(body: AlertBody, user: dict = Depends(current_user)) -> dict:
    if not (body.query.strip() or body.chamber or body.theme):
        raise HTTPException(status_code=400, detail="An alert needs words, a chamber or a subject")
    with db.session() as s:
        if (s.scalar(select(func.count()).select_from(Alert).where(Alert.user_email == user["email"])) or 0) >= 50:
            raise HTTPException(status_code=409, detail="50 alerts at most")
        a = Alert(user_email=user["email"], name=(body.name or body.query or body.theme or body.chamber)[:120],
                  query=body.query.strip()[:200],
                  filters=json.dumps({k: getattr(body, k) for k in ("chamber", "theme", "year", "city")},
                                     ensure_ascii=False),
                  lang=body.lang if body.lang in ("ar", "fr") else "ar",
                  frequency=body.frequency if body.frequency in ("daily", "weekly") else "weekly",
                  # start from now: the first email brings what is new, not the whole corpus
                  last_sent_at=dt.datetime.utcnow())
        s.add(a)
        s.commit()
        return _alert(a)


@router.patch("/api/me/alerts/{alert_id}")
def alert_update(alert_id: int, body: AlertBody, user: dict = Depends(current_user)) -> dict:
    with db.session() as s:
        a = s.get(Alert, alert_id)
        if not a or a.user_email != user["email"]:
            raise HTTPException(status_code=404, detail="Not found")
        if body.active is not None:
            a.active = int(body.active)
        if body.frequency in ("daily", "weekly"):
            a.frequency = body.frequency
        if body.name:
            a.name = body.name[:120]
        s.commit()
        return _alert(a)


@router.delete("/api/me/alerts/{alert_id}")
def alert_delete(alert_id: int, user: dict = Depends(current_user)) -> dict:
    with db.session() as s:
        a = s.get(Alert, alert_id)
        if a and a.user_email == user["email"]:
            s.delete(a)
            s.commit()
    return {"ok": True}


# -- history --
@router.get("/api/me/history")
def history(user: dict = Depends(current_user)) -> list:
    rows = db.activity_for(user["email"], limit=200)
    views = [r["decision_id"] for r in rows if r["action"] in ("view", "download") and r["decision_id"]]
    with db.session() as s:
        decs = {d.doc_id: d for d in s.scalars(select(db.Decision).where(db.Decision.doc_id.in_(views)))} if views else {}
    out = []
    for r in rows:
        if r["action"] not in ("view", "search", "download"):
            continue
        item = {"ts": r["ts"], "action": r["action"]}
        if r["action"] == "search":
            item["query"] = r["decision_id"]
        else:
            d = decs.get(r["decision_id"])
            if not d or d.state != "published":
                continue
            item.update(doc_id=d.doc_id, decision_no=d.decision_no, date=d.decision_date,
                        chamber=d.category)
        out.append(item)
    return out


class BeatBody(BaseModel):
    seconds: int
    lang: str = "ar"


@router.post("/api/decisions/{doc_id}/reading")
def reading_beat(doc_id: str, body: BeatBody, user: dict = Depends(current_user)) -> dict:
    """Add a beat of reading time. At most a minute per beat, whatever the page claims."""
    seconds = max(0, min(int(body.seconds), 60))
    if seconds == 0 or user["role"] in STAFF:
        return {"ok": True}
    lang = body.lang if body.lang in ("ar", "fr") else "ar"
    today = dt.date.today()
    with db.session() as s:
        row = s.scalar(select(ReadingTime).where(ReadingTime.actor == user["email"], ReadingTime.doc_id == doc_id,
                                                 ReadingTime.day == today, ReadingTime.lang == lang))
        if row is None:
            if not s.get(db.Decision, doc_id):
                raise HTTPException(status_code=404, detail="Not found")
            row = ReadingTime(actor=user["email"], doc_id=doc_id, day=today, lang=lang, seconds=0)
            s.add(row)
        row.seconds += seconds
        s.commit()
    return {"ok": True}


@router.get("/api/me/downloads")
def my_downloads(user: dict = Depends(current_user)) -> dict:
    import billing
    return {"used_today": downloads_today(user["email"]),
            "quota": None if user["role"] in STAFF else billing.daily_downloads(user.get("org_id"))}
