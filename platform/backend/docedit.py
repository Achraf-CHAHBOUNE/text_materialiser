"""Editing a delivered ruling: the text, the Word file, and what kind of change it is.

No database here -- only the rules, so they can be tested on their own.

The text a reader sees IS the Word file: an imported ruling's text is its paragraphs
joined by newlines, one line per paragraph (checked on 600 delivered files: 600 of
600 match, and a page break reads as an empty line). So an edit is applied to the
file itself, paragraph by paragraph, and the text is then read back out of the saved
file. Untouched paragraphs keep their exact formatting -- title, centred headings,
page breaks -- and the page and the download can never disagree.

Changes come in two kinds, and the difference is the whole safety model:

  - removal: every character left in the new text was already visible in that spot.
    Hiding a name is a removal. It cannot publish anything new, so it goes live.
  - addition: anything else -- a corrected word, a restored name, a pasted sentence.
    It could put a real name back, which no pattern check can recognise, so it waits
    as a draft until an admin approves it.
"""
from __future__ import annotations

import copy
import difflib
import io
import re
import zipfile
from typing import List, Optional, Tuple

from docx import Document
from docx.oxml.ns import qn

TOKEN = "XXXXXXX"

# ---------------------------------------------------------------- matching a name
# Close to the pipeline's redactor (pipeline/anonymizer/core/redactor.py): alef forms
# are interchangeable, and tatweel/diacritics may sit between letters. Two deliberate
# differences, because here a person picks one name to hide:
#   - ي and ى stay distinct. With them equal, hiding the name "علي" also hid "على"
#     ("on"), one of the commonest words in a ruling.
#   - the gap between letters never crosses a line. A line is a paragraph of the
#     file, and a match that swallowed a newline would merge two paragraphs.
_EQUIV = {c: "اأإآٱ" for c in "اأإآٱ"}
_GAP = "[ \\t\u00a0ـً-ٰٟ]*"
_NOISE = re.compile("[\\sـً-ٰٟ]")
_LETTER = "[ء-ؿف-يپچکگی]"
# Always anchored to whole words here, even for long values: an admin hiding "علي"
# everywhere must not turn "عليه" into "XXXXXXXه". A one- or two-letter clitic in
# front ("وعلي", "فبعلي") still matches, as in the pipeline.
_WORD_START = (f"(?:(?<!{_LETTER})|(?<=(?<!{_LETTER})[وفبلك])"
               f"|(?<=(?<!{_LETTER})[وف][بلك]))")
_WORD_END = f"(?!{_LETTER})"
# Fewer significant characters than this would hit fragments of ordinary words.
MIN_HIDE_ALL_CHARS = 3


class EditError(ValueError):
    """The edit cannot be applied as asked; the message says why."""


def significant(value: str) -> str:
    return _NOISE.sub("", value or "")


def name_pattern(value: str) -> re.Pattern:
    sig = significant(value)
    if len(sig) < MIN_HIDE_ALL_CHARS:
        raise EditError(f"at least {MIN_HIDE_ALL_CHARS} letters are needed to hide everywhere")
    if TOKEN in value:
        raise EditError("that text is already hidden")
    body = _GAP.join(f"[{_EQUIV[c]}]" if c in _EQUIV else re.escape(c) for c in sig)
    return re.compile(f"{_WORD_START}(?:{body}){_WORD_END}")


def _replace_spans(text: str, spans: List[Tuple[int, int]]) -> str:
    """Put the token over each span. A span never crosses a line (see _GAP), but a
    selection can: then each line's part is hidden on its own, keeping the lines."""
    out, last = [], 0
    for start, end in sorted(spans):
        if start < last:
            continue                                   # overlapping: already covered
        out.append(text[last:start])
        out.append("\n".join(TOKEN if part.strip() else part
                             for part in text[start:end].split("\n")))
        last = end
    out.append(text[last:])
    return "".join(out)


def find_all(text: str, value: str) -> List[Tuple[int, int]]:
    """Every match of `value` as a whole word, variant-tolerant."""
    rx = name_pattern(value)
    return [m.span() for m in rx.finditer(text)]


def hide_everywhere(text: str, value: str) -> Tuple[str, int]:
    spans = find_all(text, value)
    return _replace_spans(text, spans), len(spans)


def hide_occurrence(text: str, value: str, occurrence: int) -> str:
    """Hide the n-th exact occurrence of the selected text (0-based, non-overlapping,
    counted the way the browser counts them with indexOf)."""
    if not value or value != value.strip():
        raise EditError("select the text itself, without surrounding spaces")
    if not significant(value).replace(TOKEN, ""):
        raise EditError("that text is already hidden")
    pos, seen = 0, 0
    while True:
        at = text.find(value, pos)
        if at < 0:
            raise EditError("the selected text is no longer in this ruling -- reload it")
        if seen == occurrence:
            return _replace_spans(text, [(at, at + len(value))])
        seen, pos = seen + 1, at + len(value)


def snippets(text: str, spans: List[Tuple[int, int]], width: int = 40, limit: int = 5) -> List[str]:
    out = []
    for start, end in spans[:limit]:
        a, b = max(0, start - width), min(len(text), end + width)
        out.append((("…" if a else "") + text[a:start] + "«" + text[start:end] + "»"
                    + text[end:b] + ("…" if b < len(text) else "")).replace("\n", " "))
    return out


# ---------------------------------------------------------------- kind of change
def _words(s: str) -> List[str]:
    return s.split()


def is_removal(old: str, new: str) -> bool:
    """True if `new` shows nothing that `old` did not already show in that place.

    Compared line by line, then word by word inside a changed stretch. A changed word
    passes only if, once its tokens are taken out, what is left was already part of
    the words it replaces: "محمد،" -> "XXXXXXX،" passes, "XXXXXXX" -> "محمد" does not,
    and neither does any word typed where there was none.
    """
    if old == new:
        return True
    a, b = old.split("\n"), new.split("\n")
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if tag == "equal":
            continue
        ow, nw = _words(" ".join(a[i1:i2])), _words(" ".join(b[j1:j2]))
        wm = difflib.SequenceMatcher(None, ow, nw, autojunk=False)
        for wtag, w1, w2, v1, v2 in wm.get_opcodes():
            if wtag in ("equal", "delete"):
                continue
            was = " ".join(ow[w1:w2])
            for word in nw[v1:v2]:
                rest = word.replace(TOKEN, "")
                if rest and rest not in was:
                    return False
    return True


# ---------------------------------------------------------------- showing a change
def diff(old: str, new: str, context: int = 6) -> List[dict]:
    """Changed stretches, word-level, with a little unchanged text around each."""
    a, b = old.split("\n"), new.split("\n")
    hunks: List[dict] = []
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if tag == "equal":
            continue
        ot, nt = "\n".join(a[i1:i2]), "\n".join(b[j1:j2])
        ow, nw = re.split(r"(\s+)", ot), re.split(r"(\s+)", nt)
        segs: List[dict] = []

        def push(op: str, text: str) -> None:
            if not text:
                return
            if segs and segs[-1]["op"] == op:
                segs[-1]["text"] += text
            else:
                segs.append({"op": op, "text": text})

        for wtag, w1, w2, v1, v2 in difflib.SequenceMatcher(None, ow, nw, autojunk=False).get_opcodes():
            if wtag == "equal":
                chunk = "".join(ow[w1:w2])
                if len(chunk) > 2 * context * 8:       # long unchanged run: keep the ends
                    chunk = chunk[:context * 8] + " … " + chunk[-context * 8:]
                push("equal", chunk)
            else:
                push("del", "".join(ow[w1:w2]))
                push("ins", "".join(nw[v1:v2]))
        hunks.append({"line": i1 + 1, "segments": segs})
    return hunks


def summarize(old: str, new: str) -> str:
    hidden = new.count(TOKEN) - old.count(TOKEN)
    changed = sum(1 for h in diff(old, new) for s in h["segments"] if s["op"] != "equal")
    if is_removal(old, new) and hidden > 0:
        return f"{hidden} hidden"
    return f"{changed} change(s)"


# ---------------------------------------------------------------- the Word file
_XML_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]")


def _holds_break(p) -> bool:
    return bool(p._p.findall(".//" + qn("w:br")))


def _text_runs(p) -> list:
    return [r for r in p.runs if not r._r.findall(qn("w:br"))]


def _body_template(paragraphs) -> Optional[object]:
    """The paragraph a new line should look like: the longest body paragraph."""
    body = [p for p in paragraphs if _text_runs(p) and not _holds_break(p)]
    return max(body, key=lambda p: len(p.text), default=None)


def _set_text(p, text: str, template) -> None:
    text = _XML_ILLEGAL.sub("", text)
    runs = _text_runs(p)
    if runs:
        runs[0].text = text
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
    elif text:
        source = _text_runs(template)[0]._r if template is not None and _text_runs(template) else None
        if source is not None:
            r = copy.deepcopy(source)
            for child in list(r):
                if child.tag != qn("w:rPr"):
                    r.remove(child)
            p._p.insert(len(p._p.findall(qn("w:pPr"))), r)   # before any break run
            p.runs[0].text = text
        else:
            p.add_run(text)


def _new_paragraph_like(anchor, template):
    source = anchor if (anchor is not None and not _holds_break(anchor)) else template
    el = copy.deepcopy(source._p) if source is not None else None
    if el is None:
        raise EditError("this file has no paragraph to copy the style from")
    for r in el.findall(qn("w:r")):
        el.remove(r)
    return el


def _save(doc) -> bytes:
    """Fixed zip timestamps, so the same content always gives the same bytes
    (the pipeline's writer does the same)."""
    raw = io.BytesIO()
    doc.save(raw)
    raw.seek(0)
    out = io.BytesIO()
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = info.external_attr
            dst.writestr(fixed, src.read(info.filename))
    return out.getvalue()


def read_text(docx: bytes) -> str:
    return "\n".join(p.text for p in Document(io.BytesIO(docx)).paragraphs)


def apply_to_docx(docx: bytes, new_text: str) -> Tuple[bytes, str]:
    """Rewrite only the paragraphs whose line changed. Returns (file, text read back).

    The text read back is what gets stored: if the file could not take an edit exactly
    (say, a deleted empty line that holds a page break, which is kept), the page shows
    what the file really says.
    """
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(docx))
    paragraphs = list(doc.paragraphs)
    old_lines = [p.text for p in paragraphs]
    new_lines = new_text.split("\n")
    template = _body_template(paragraphs)

    ops = difflib.SequenceMatcher(None, old_lines, new_lines, autojunk=False).get_opcodes()
    for tag, i1, i2, j1, j2 in ops:
        if tag == "equal":
            continue
        olds, news = paragraphs[i1:i2], new_lines[j1:j2]
        for p, line in zip(olds, news):
            _set_text(p, line, template)
        if len(news) > len(olds):
            anchor = olds[-1] if olds else (paragraphs[i1 - 1] if i1 > 0 else None)
            for line in news[len(olds):]:
                el = _new_paragraph_like(anchor, template)
                if anchor is not None:
                    anchor._p.addnext(el)
                elif paragraphs:
                    paragraphs[0]._p.addprevious(el)
                else:
                    doc.element.body.append(el)
                new_p = Paragraph(el, anchor._parent if anchor is not None else doc._body)
                _set_text(new_p, line, template)
                anchor = new_p
        for p in olds[len(news):]:
            if _holds_break(p):
                _set_text(p, "", template)                 # keep the page break
            else:
                p._p.getparent().remove(p._p)

    data = _save(doc)
    return data, read_text(data)


def retitle(docx: bytes, old_title: str, new_title: str) -> Optional[Tuple[bytes, str]]:
    """Change the title line if it still reads as the old "court — chamber" title."""
    text = read_text(docx)
    first, _, rest = text.partition("\n")
    if not new_title or first != old_title or first == new_title:
        return None
    return apply_to_docx(docx, new_title + ("\n" + rest if "\n" in text else ""))


def title_for(court: str, chamber: str) -> str:
    return " — ".join(p for p in (court, chamber) if p and p != "غير محدد")
