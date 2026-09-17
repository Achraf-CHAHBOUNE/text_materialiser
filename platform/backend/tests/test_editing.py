"""Editing a ruling from the site: the Word file follows the text, hiding goes live,
adding waits for approval, every change can be undone, and nothing undoes it by accident.

Env matches test_platform.py exactly: whichever module is imported first configures
the shared app, and both need the same settings.
"""
import io
import json
import os
import tempfile
import zipfile

_tmp = tempfile.mkdtemp()
os.environ.update(
    DB_DIR=_tmp, FILESTORE_DIR=os.path.join(_tmp, "files"),
    AUTH_SECRET="test-secret", SEED_EMAIL="admin@x.com", SEED_PASSWORD="admin-pass",
    AUTO_PUBLISH="0",
)
os.environ.pop("DATABASE_URL", None)
os.environ.pop("MINIO_ENDPOINT", None)

import pytest  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_BREAK  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import docedit  # noqa: E402
import server  # noqa: E402

server._startup()
client = TestClient(server.app)


@pytest.fixture(scope="module", autouse=True)
def _leave_no_trace():
    """Both test modules share one database (the engine is made at import). Remove
    every row and file made here, so test_platform's exact counts still hold."""
    yield
    import db
    from sqlalchemy import delete, select
    with db.session() as s:
        ids = list(s.scalars(select(db.Decision.doc_id).where(db.Decision.doc_id.like("edit_%"))))
        for doc_id in ids:
            server.STORE.delete(f"{doc_id}.docx")
        for model in (db.DecisionVersion, db.Report, db.Link):
            col = model.doc_id if hasattr(model, "doc_id") else model.parent_doc
            s.execute(delete(model).where(col.in_(ids)))
        s.execute(delete(db.Decision).where(db.Decision.doc_id.in_(ids)))
        s.execute(delete(db.ClientActivity).where(db.ClientActivity.actor == "reader@x.com"))
        s.execute(delete(db.User).where(db.User.email == "reader@x.com"))
        s.commit()

NAME = "محمد العلوي"
BODY = [
    "محكمة النقض — مدنية",
    "",
    "باسم جلالة الملك",
    f"بين السيد {NAME} الساكن بالرباط والمطلوب على الخصوص في النقض",
    "وحيث إن الطاعن علي يعيب على القرار خرق القانون، وعليه فإن الوسيلة غير مقبولة",
    "",                                   # holds the page break
    "لهذه الأسباب",
    f"قضت محكمة النقض برفض الطلب وتحميل {NAME} الصائر",
]


def _ruling_docx(lines=BODY) -> bytes:
    """Shaped like the pipeline's writer: a bold title, a page break paragraph."""
    doc = Document()
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        if i == 5:
            p.add_run().add_break(WD_BREAK.PAGE)
            continue
        if line:
            run = p.add_run(line)
            run.bold = i in (0, 2, 6)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _token(email, pw):
    return client.post("/api/auth/login", json={"email": email, "password": pw}).json()["token"]


def _h(tok):
    return {"Authorization": f"Bearer {tok}"}


ADMIN = {}


def admin():
    if "h" not in ADMIN:
        ADMIN["h"] = _h(_token("admin@x.com", "admin-pass"))
    return ADMIN["h"]


def reader():
    if "r" not in ADMIN:
        client.post("/api/admin/clients", headers=admin(),
                    json={"email": "reader@x.com", "password": "reader-pass"})
        ADMIN["r"] = _h(_token("reader@x.com", "reader-pass"))
    return ADMIN["r"]


def _import(doc_id: str, data: bytes | None = None) -> dict:
    rec = {"doc_id": doc_id, "anonymized_file": f"{doc_id}.docx", "format": "docx",
           "fields": {"رقم القرار": {"value": "77"}, "تاريخ القرار": {"value": "12/03/2021"}},
           "category": {"value": "مدنية"}, "level": "نقض", "court": "محكمة النقض",
           "pii": {"removed_count": 0}, "review": "ok"}
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("records.json", json.dumps([rec], ensure_ascii=False))
        z.writestr(f"{doc_id}.docx", data or _ruling_docx())
    r = client.post("/api/admin/import", headers=admin(),
                    files={"file": ("b.zip", buf.getvalue(), "application/zip")})
    assert r.status_code == 200, r.text
    client.post(f"/api/admin/decisions/{doc_id}/state", headers=admin(), json={"state": "published"})
    return r.json()


def _live(doc_id):
    return client.get(f"/api/admin/decisions/{doc_id}", headers=admin()).json()


def _file(doc_id) -> bytes:
    return server.STORE.get(_live(doc_id)["file_name"])


# ------------------------------------------------------------------ the Word file
def test_an_edit_changes_only_its_paragraph_and_keeps_the_layout():
    original = _ruling_docx()
    lines = docedit.read_text(original).split("\n")
    lines[3] = lines[3].replace(NAME, "XXXXXXX")
    data, text = docedit.apply_to_docx(original, "\n".join(lines))

    assert text == "\n".join(lines)
    before, after = Document(io.BytesIO(original)), Document(io.BytesIO(data))
    assert len(after.paragraphs) == len(before.paragraphs)
    for i, (a, b) in enumerate(zip(before.paragraphs, after.paragraphs)):
        if i != 3:                               # every other paragraph byte-for-byte
            assert a._p.xml == b._p.xml, i
    assert after.paragraphs[0].runs[0].bold      # title formatting kept
    assert len(after.element.body.findall(".//" + qn("w:br"))) == 1   # page break kept


def test_added_and_removed_lines_land_in_the_right_place():
    original = _ruling_docx()
    lines = docedit.read_text(original).split("\n")
    new = lines[:4] + ["سطر جديد"] + lines[4:7] + lines[8:]      # insert one, drop the last
    data, text = docedit.apply_to_docx(original, "\n".join(new))
    assert text == "\n".join(new)
    assert len(Document(io.BytesIO(data)).element.body.findall(".//" + qn("w:br"))) == 1


def test_the_same_edit_gives_the_same_bytes():
    original = _ruling_docx()
    new = docedit.read_text(original).replace(NAME, "XXXXXXX")
    assert docedit.apply_to_docx(original, new)[0] == docedit.apply_to_docx(original, new)[0]


# ------------------------------------------------------------------ what kind of change
def test_hiding_is_a_removal_and_anything_that_shows_more_is_not():
    old = f"بين {NAME}، الساكن بالرباط"
    assert docedit.is_removal(old, old.replace(NAME, "XXXXXXX"))
    assert docedit.is_removal(old, "بين XXXXXXX، الساكن")                 # deleting text
    assert not docedit.is_removal(old.replace(NAME, "XXXXXXX"), old)      # un-hiding
    assert not docedit.is_removal(old, old + " أحمد")                     # typing a word
    assert not docedit.is_removal(old, old + "\nسطر جديد")                 # adding a line
    assert not docedit.is_removal(old, "بين الساكن محمد بالرباط")          # moving a word


def test_hide_everywhere_takes_whole_words_only():
    text = "\n".join(BODY)
    new, n = docedit.hide_everywhere(text, "علي")
    assert n == 1
    assert "على القرار" in new and "على الخصوص" in new     # على is not the name علي
    assert "وعليه" in new                                  # nor is part of a longer word
    new, n = docedit.hide_everywhere(text, NAME)
    assert n == 2 and NAME not in new
    new, n = docedit.hide_everywhere("وبمحمد العلوي حضر", NAME)
    assert n == 1                                          # behind clitics, still a match


def test_two_letters_are_not_enough_to_hide_everywhere():
    try:
        docedit.hide_everywhere("\n".join(BODY), "عل")
    except docedit.EditError:
        return
    raise AssertionError("a two-letter hide-everywhere must be refused")


# ------------------------------------------------------------------ through the API
def test_hiding_a_name_goes_live_and_rewrites_the_file():
    _import("edit_hide")
    d = _live("edit_hide")
    r = client.post("/api/admin/decisions/edit_hide/hide", headers=admin(),
                    json={"value": NAME, "occurrence": 1, "base_version": d["version"]})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["applied"] and out["version"] == 2
    body = out["decision"]["body_text"]
    assert body.count(NAME) == 1 and body.index(NAME) < body.index("XXXXXXX")   # the 2nd one
    assert docedit.read_text(_file("edit_hide")) == body                         # file == page
    assert [v["version"] for v in out["history"]["versions"]] == [2, 1]

    seen = client.get("/api/decisions/edit_hide", headers=reader()).json()["body_text"]
    assert seen == body


def test_adding_text_waits_as_a_draft_until_approved():
    _import("edit_draft")
    d = _live("edit_draft")
    new = d["body_text"].replace("لهذه الأسباب", "لهذه الأسباب المذكورة")
    r = client.post("/api/admin/decisions/edit_draft/text", headers=admin(),
                    json={"text": new, "base_version": d["version"]}).json()
    assert r["applied"] is False and r["draft_id"]
    assert _live("edit_draft")["body_text"] == d["body_text"]            # nothing live yet

    ok = client.post(f"/api/admin/drafts/{r['draft_id']}/approve", headers=admin())
    assert ok.status_code == 200, ok.text
    assert _live("edit_draft")["body_text"] == new
    assert docedit.read_text(_file("edit_draft")) == new


def test_an_edit_made_on_an_old_version_is_refused():
    _import("edit_stale")
    d = _live("edit_stale")
    draft = client.post("/api/admin/decisions/edit_stale/text", headers=admin(), json={
        "text": d["body_text"] + " إضافة", "base_version": d["version"]}).json()["draft_id"]
    client.post("/api/admin/decisions/edit_stale/hide", headers=admin(),
                json={"value": NAME, "base_version": d["version"]})
    again = client.post("/api/admin/decisions/edit_stale/hide", headers=admin(),
                        json={"value": NAME, "base_version": d["version"]})
    assert again.status_code == 409                     # someone saved in between
    assert client.post(f"/api/admin/drafts/{draft}/approve", headers=admin()).status_code == 409


def test_restoring_brings_back_the_exact_file_but_only_once_approved():
    original = _ruling_docx()
    _import("edit_restore", original)
    d = _live("edit_restore")
    stored_original = _file("edit_restore")
    client.post("/api/admin/decisions/edit_restore/hide", headers=admin(),
                json={"value": NAME, "everywhere": True, "base_version": d["version"]})
    live = _live("edit_restore")
    assert NAME not in live["body_text"]

    r = client.post("/api/admin/decisions/edit_restore/restore", headers=admin(),
                    json={"version": 1, "base_version": live["version"]}).json()
    assert r["applied"] is False                        # it would show the name again
    client.post(f"/api/admin/drafts/{r['draft_id']}/approve", headers=admin())
    assert _file("edit_restore") == stored_original


def test_purging_history_deletes_the_last_copy_of_a_hidden_name():
    _import("edit_purge")
    d = _live("edit_purge")
    client.post("/api/admin/decisions/edit_purge/hide", headers=admin(),
                json={"value": NAME, "everywhere": True, "base_version": d["version"]})
    live = _live("edit_purge")
    out = client.post("/api/admin/decisions/edit_purge/purge-history", headers=admin()).json()
    assert out["deleted"] == 1 and [v["version"] for v in out["history"]["versions"]] == [2]
    gone = client.post("/api/admin/decisions/edit_purge/restore", headers=admin(),
                       json={"version": 1, "base_version": live["version"]})
    assert gone.status_code == 404


def test_a_reimport_never_undoes_a_hand_edit():
    _import("edit_reimport")
    d = _live("edit_reimport")
    client.post("/api/admin/decisions/edit_reimport/hide", headers=admin(),
                json={"value": NAME, "everywhere": True, "base_version": d["version"]})
    edited_file = _file("edit_reimport")

    out = _import("edit_reimport", _ruling_docx())       # the pipeline's version again
    assert out["kept_edited"] == ["edit_reimport"]
    assert NAME not in _live("edit_reimport")["body_text"]
    assert _file("edit_reimport") == edited_file


def test_an_edit_that_adds_personal_data_is_refused_outright():
    _import("edit_gate")
    d = _live("edit_gate")
    r = client.post("/api/admin/decisions/edit_gate/text", headers=admin(), json={
        "text": d["body_text"] + " الهاتف 0612345678", "base_version": d["version"]})
    assert r.status_code == 422
    assert _live("edit_gate")["version"] == d["version"]
    assert client.get("/api/admin/decisions/edit_gate/history", headers=admin()).json()["drafts"] == []


def test_changing_the_chamber_retitles_the_file_and_resorts_the_listing():
    _import("edit_meta")
    d = _live("edit_meta")
    bad = client.patch("/api/admin/decisions/edit_meta", headers=admin(),
                       json={"changes": {"decision_date": "2021-03-12"}})
    assert bad.status_code == 422
    r = client.patch("/api/admin/decisions/edit_meta", headers=admin(), json={"changes": {
        "category": "جنائية", "city": "فاس", "decision_date": "05/11/2019"}})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["city"] == "فاس" and out["year"] == "2019" and out["version"] == d["version"] + 2
    assert docedit.read_text(_file("edit_meta")).split("\n")[0] == "محكمة النقض — جنائية"
    row = next(x for x in client.get("/api/browse/rulings?chamber=جنائية", headers=reader())
               .json()["items"] if x["doc_id"] == "edit_meta")
    assert row["date"] == "05/11/2019"


def test_a_reader_reports_a_problem_and_the_quote_is_dropped_once_handled():
    _import("edit_report")
    assert client.post("/api/admin/decisions/edit_report/hide", headers=reader(),
                       json={"value": NAME, "base_version": 0}).status_code == 403
    r = client.post("/api/decisions/edit_report/reports", headers=reader(),
                    json={"quote": NAME, "note": "اسم ظاهر"})
    assert r.status_code == 200, r.text
    open_ = client.get("/api/admin/decisions/edit_report/history", headers=admin()).json()["reports"]
    assert open_ and open_[0]["quote"] == NAME
    client.post(f"/api/admin/reports/{r.json()['id']}", headers=admin(), json={"status": "resolved"})
    done = client.get("/api/admin/reports?status=resolved", headers=admin()).json()
    assert next(x for x in done if x["id"] == r.json()["id"])["quote"] == ""


def test_export_carries_every_edited_ruling_and_its_file():
    _import("edit_export")
    d = _live("edit_export")
    client.post("/api/admin/decisions/edit_export/hide", headers=admin(),
                json={"value": NAME, "everywhere": True, "base_version": d["version"]})
    r = client.get("/api/admin/edits/export", headers=admin())
    assert r.status_code == 200
    z = zipfile.ZipFile(io.BytesIO(r.content))
    edits = {e["doc_id"]: e for e in json.loads(z.read("edits.json"))}
    assert "edit_export" in edits
    assert docedit.read_text(z.read("files/edit_export.docx")) == _live("edit_export")["body_text"]


def test_hide_everywhere_can_be_previewed_without_changing_anything():
    _import("edit_preview")
    d = _live("edit_preview")
    r = client.post("/api/admin/decisions/edit_preview/hide", headers=admin(), json={
        "value": NAME, "everywhere": True, "dry_run": True, "base_version": d["version"]}).json()
    assert r["count"] == 2 and all("«" in s for s in r["snippets"])
    assert _live("edit_preview")["version"] == d["version"]
