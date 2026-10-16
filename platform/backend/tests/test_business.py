"""The business side, followed the way it is used: a demo request becomes a customer,
its people are invited and read within what they pay for, and the team is paid.

Env matches the other test modules: whichever is imported first configures the app.
"""
import datetime as dt
import io
import os
import tempfile

_tmp = tempfile.mkdtemp()
os.environ.update(
    DB_DIR=_tmp, FILESTORE_DIR=os.path.join(_tmp, "files"),
    AUTH_SECRET="test-secret", SEED_EMAIL="admin@x.com", SEED_PASSWORD="admin-pass",
    AUTO_PUBLISH="0",
)
os.environ.pop("DATABASE_URL", None)
os.environ.pop("SMTP_HOST", None)

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from sqlalchemy import select  # noqa: E402

import accounts  # noqa: E402
import auth as authmod  # noqa: E402
import billing  # noqa: E402
import db  # noqa: E402
import library  # noqa: E402
import mailer  # noqa: E402
import server  # noqa: E402
import storage  # noqa: E402
import textnorm  # noqa: E402
import worker  # noqa: E402

server._startup()
db.init_db()                    # the worker's own table
client = TestClient(server.app)
PW = "a-long-password"


def _login(email, pw=PW, code=""):
    r = client.post("/api/auth/login", json={"email": email, "password": pw, "code": code})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['token']}"}


def _staff(email, role):
    if not db.get_user(email):
        db.upsert_user(email, authmod.hash_password(PW), role=role)
    return _login(email)


ADMIN = None


def admin():
    global ADMIN
    ADMIN = ADMIN or _login("admin@x.com", "admin-pass")
    return ADMIN


def _last_mail(to, template):
    with db.session() as s:
        m = s.scalar(select(mailer.EmailOutbox).where(mailer.EmailOutbox.to_addr == to,
                                                      mailer.EmailOutbox.template == template)
                     .order_by(mailer.EmailOutbox.id.desc()))
        return (m.subject, m.body) if m else None


def _token_from(body):
    return body.split("/invite/")[1].split()[0] if "/invite/" in body else body.split("/reset/")[1].split()[0]


def _ruling(doc_id, created=None):
    from docx import Document
    d = Document()
    d.add_paragraph("وحيث إن عقد الكراء صحيح")
    b = io.BytesIO()
    d.save(b)
    storage.STORE.put(f"{doc_id}.docx", b.getvalue())
    with db.session() as s:
        s.merge(db.Decision(doc_id=doc_id, category="مدنية", court="محكمة النقض",
                            decision_no=doc_id[-3:], decision_date="01/02/2020", year="2020",
                            state="published", file_name=f"{doc_id}.docx",
                            body_text="وحيث إن عقد الكراء صحيح",
                            search_text=textnorm.normalize("وحيث إن عقد الكراء صحيح"),
                            sort_key=f"20200201{doc_id[-3:]}",
                            created_at=created or dt.datetime.utcnow()))
        s.commit()


@pytest.fixture(scope="module")
def customer():
    """A client the way they come in Morocco: the transfer, the contract, then the
    account -- created by the admin in one step; the holder through the invitation."""
    sales = _staff("sales@x.com", "sales")
    out = client.post("/api/crm/clients", headers=admin(), json={
        "org": {"name": "Cabinet Test", "kind": "firm", "email": "contact@firm.ma", "language": "fr"},
        "contract": {"plan_code": "firm", "seats": 2, "amount_ht": 1200000},
        "transfer": {"amount": 1440000, "bank_ref": "VIR-0001"},
        "accounts": [{"email": "boss@firm.ma", "role": "client_admin"}]}).json()
    assert out["active"] is True
    org_id = out["org_id"]
    body = _last_mail("boss@firm.ma", "invitation")[1]
    assert "/fr/invite/" in body
    token = _token_from(body)
    assert client.get(f"/api/invitations/{token}").json()["org"] == "Cabinet Test"
    assert client.post(f"/api/invitations/{token}", json={"password": PW, "full_name": "Boss"}).status_code == 200
    # the link works once
    assert client.post(f"/api/invitations/{token}", json={"password": PW}).status_code == 404
    return {"org_id": org_id, "sales": sales, "boss": _login("boss@firm.ma")}


def test_the_public_form_lands_in_the_inbox_and_is_guarded():
    r = client.post("/api/public/demo", json={"full_name": "Maître X", "email": "x@cab.ma", "language": "fr"})
    assert r.status_code == 200
    assert _last_mail("x@cab.ma", "demo_received")[0].startswith("Nous avons")
    assert any(l["email"] == "x@cab.ma" for l in client.get("/api/crm/leads", headers=admin()).json())
    bot = client.post("/api/public/demo", json={"full_name": "x", "email": "bot@x.ma", "website": "spam"})
    assert bot.status_code == 200
    with db.session() as s:
        import crm
        assert s.scalar(select(crm.Lead).where(crm.Lead.email == "bot@x.ma")) is None


def test_a_paid_contract_opens_the_accounts(customer):
    acc = client.get("/api/account", headers=customer["boss"]).json()
    assert acc["user"]["role"] == "client_admin" and acc["org"]["name"] == "Cabinet Test"
    assert acc["subscription"]["status"] == "active" and acc["access"] == "ok"
    assert acc["seats"] == {"total": 2, "used": 1}
    assert client.get("/api/browse/courts", headers=customer["boss"]).status_code == 200


def test_no_transfer_no_access_until_it_comes(customer):
    out = client.post("/api/crm/clients", headers=admin(), json={
        "org": {"name": "Pas encore payé"}, "contract": {"plan_code": "individual", "amount_ht": 500000},
        "accounts": [{"email": "unpaid@cab.ma"}]}).json()
    assert out["active"] is False
    token = _token_from(_last_mail("unpaid@cab.ma", "invitation")[1])
    client.post(f"/api/invitations/{token}", json={"password": PW})
    h = _login("unpaid@cab.ma")
    assert client.get("/api/browse/courts", headers=h).status_code == 402      # account made, closed
    c = client.get(f"/api/crm/orgs/{out['org_id']}", headers=admin()).json()["contracts"][0]
    assert c["status"] == "awaiting_payment" and c["balance"] == 600000          # VAT 20% included
    client.post(f"/api/crm/contracts/{c['id']}/transfers", headers=admin(), json={"amount": 600000, "bank_ref": "VIR-2"})
    assert client.get("/api/browse/courts", headers=h).status_code == 200      # open the moment it is paid
    c = client.get(f"/api/crm/orgs/{out['org_id']}", headers=admin()).json()["contracts"][0]
    assert c["status"] == "active" and c["balance"] == 0


def test_more_accounts_than_the_contract_is_refused_before_anything_is_made():
    r = client.post("/api/crm/clients", headers=admin(), json={
        "org": {"name": "Trop de comptes"}, "contract": {"plan_code": "individual", "seats": 1},
        "accounts": [{"email": "a1@trop.ma"}, {"email": "a2@trop.ma"}]})
    assert r.status_code == 400
    assert not any(o["name"] == "Trop de comptes" for o in client.get("/api/crm/orgs", headers=admin()).json())


def test_a_transfer_can_become_an_invoice_once(customer):
    c = client.get(f"/api/crm/orgs/{customer['org_id']}", headers=admin()).json()["contracts"][-1]
    t = c["transfers"][0]
    first = client.post(f"/api/crm/transfers/{t['id']}/invoice", headers=admin()).json()
    again = client.post(f"/api/crm/transfers/{t['id']}/invoice", headers=admin()).json()
    assert first["invoice_id"] == again["invoice_id"]
    inv = client.get(f"/api/invoices/{first['invoice_id']}", headers=admin()).json()
    assert inv["status"] == "paid" and inv["amount_ttc"] == 1440000
    # counted once: the transfer, not the transfer and its invoice
    org = next(o for o in client.get("/api/crm/orgs", headers=admin()).json() if o["id"] == customer["org_id"])
    assert org["paid"] >= 1440000 and org["paid"] < 2 * 1440000


def test_seats_limit_invitations(customer):
    r = client.post("/api/account/team/invite", headers=customer["boss"],
                    json={"email": "clerk@firm.ma", "role": "client"})
    assert r.status_code == 200
    full = client.post("/api/account/team/invite", headers=customer["boss"],
                       json={"email": "third@firm.ma", "role": "client"})
    assert full.status_code == 409 and full.json()["detail"] == {"seats": 2}


def test_a_reader_cannot_manage_the_team_or_reach_staff_tools(customer):
    token = _token_from(_last_mail("clerk@firm.ma", "invitation")[1])
    client.post(f"/api/invitations/{token}", json={"password": PW})
    clerk = _login("clerk@firm.ma")
    assert client.post("/api/account/team/invite", headers=clerk,
                       json={"email": "z@firm.ma"}).status_code == 403
    assert client.get("/api/crm/orgs", headers=clerk).status_code == 403
    assert client.get("/api/admin/decisions", headers=clerk).status_code == 403


def test_one_seat_is_one_person_at_a_time(customer):
    first = _login("clerk@firm.ma")
    second = _login("clerk@firm.ma")
    assert client.get("/api/account", headers=first).status_code == 401
    assert client.get("/api/account", headers=second).status_code == 200


def test_staff_roles_are_kept_apart():
    editor = _staff("editor@x.com", "editor")
    sales = _staff("sales2@x.com", "sales")
    assert client.get("/api/crm/orgs", headers=editor).status_code == 403
    assert client.get("/api/admin/decisions", headers=sales).status_code == 403
    assert client.get("/api/admin/decisions", headers=editor).status_code == 200
    assert client.get("/api/admin/staff", headers=sales).status_code == 403


def test_a_lapsed_subscription_closes_the_rulings_not_the_account(customer):
    with db.session() as s:
        for x in s.scalars(select(billing.Subscription).where(
                billing.Subscription.org_id == customer["org_id"])):
            x.ends_on = dt.date.today() - dt.timedelta(days=1)
        s.commit()
    worker.expire_subscriptions(dt.date.today())
    boss = customer["boss"]
    r = client.get("/api/browse/courts", headers=boss)
    assert r.status_code == 402 and r.json()["detail"] == {"subscription": "expired"}
    assert client.get("/api/account", headers=boss).json()["access"] == "expired"
    assert client.get("/api/support", headers=boss).status_code == 200
    # renewed: a new contract, a new transfer -- open again, and its seats with it
    today = dt.date.today()
    c = client.post(f"/api/crm/orgs/{customer['org_id']}/contracts", headers=admin(), json={
        "plan_code": "firm", "seats": 5, "amount_ht": 1200000, "starts_on": today.isoformat(),
        "ends_on": (today + dt.timedelta(days=364)).isoformat()}).json()
    assert client.get("/api/browse/courts", headers=boss).status_code == 402        # not before the money
    client.post(f"/api/crm/contracts/{c['id']}/transfers", headers=admin(), json={"amount": 1440000})
    assert client.get("/api/browse/courts", headers=boss).status_code == 200
    assert client.get("/api/account", headers=boss).json()["seats"]["total"] == 5


def test_invoices_from_transfers_are_numbered_without_gaps(customer):
    oid, sales = customer["org_id"], customer["sales"]
    c = client.post(f"/api/crm/orgs/{oid}/contracts", headers=sales,
                    json={"plan_code": "firm", "amount_ht": 149900})
    c = next(x for x in client.get(f"/api/crm/orgs/{oid}", headers=sales).json()["contracts"] if x["id"] == c.json()["id"])
    assert c["amount_ttc"] == 179880                                        # 20% VAT, exact
    t1 = client.post(f"/api/crm/contracts/{c['id']}/transfers", headers=sales, json={"amount": 79880}).json()
    t2 = client.post(f"/api/crm/contracts/{c['id']}/transfers", headers=sales, json={"amount": 100000}).json()
    a = client.post(f"/api/crm/transfers/{t1['id']}/invoice", headers=sales).json()
    b = client.post(f"/api/crm/transfers/{t2['id']}/invoice", headers=sales).json()
    year = dt.date.today().year
    assert a["number"].startswith(f"F-{year}-")
    assert int(b["number"].rsplit("-", 1)[1]) == int(a["number"].rsplit("-", 1)[1]) + 1
    assert client.get(f"/api/invoices/{a['invoice_id']}", headers=admin()).json()["status"] == "paid"
    # the account holder sees the firm's invoices, a reader none
    mine = client.get("/api/account/invoices", headers=customer["boss"]).json()
    assert {i["number"] for i in mine} >= {a["number"], b["number"]}
    assert client.get(f"/api/invoices/{a['invoice_id']}", headers=_login("clerk@firm.ma")).status_code == 404


def test_folders_are_private_unless_shared_with_the_firm(customer):
    _ruling("lib_001")
    boss, clerk = customer["boss"], _login("clerk@firm.ma")
    mine = client.post("/api/me/folders", headers=boss, json={"name": "Baux"}).json()
    shared = client.post("/api/me/folders", headers=boss, json={"name": "Cabinet", "shared": True}).json()
    client.post(f"/api/me/folders/{shared['id']}/items", headers=boss, json={"doc_id": "lib_001", "note": "voir"})
    seen = {f["name"] for f in client.get("/api/me/folders", headers=clerk).json()}
    assert seen == {"Cabinet"}
    detail = client.get(f"/api/me/folders/{shared['id']}", headers=clerk).json()
    assert [i["doc_id"] for i in detail["items"]] == ["lib_001"] and detail["items"][0]["note"] == "voir"
    assert client.get(f"/api/me/folders/{mine['id']}", headers=clerk).status_code == 404


def test_downloads_are_counted_limited_and_marked(customer):
    _ruling("dl_001")
    boss = customer["boss"]
    r = client.get("/api/decisions/dl_001/file", headers=boss)
    assert r.status_code == 200
    from docx import Document
    footer = Document(io.BytesIO(r.content)).sections[0].footer.paragraphs[-1].text
    assert "boss@firm.ma" in footer and "réf." in footer
    org_id = customer["org_id"]
    quota = billing.daily_downloads(org_id)
    with db.session() as s:
        for i in range(quota):
            s.add(library.Download(ref=f"T{i}", user_email="boss@firm.ma", org_id=org_id, doc_id="dl_001"))
        s.commit()
    r = client.get("/api/decisions/dl_001/file", headers=boss)
    assert r.status_code == 429 and r.json()["detail"]["quota"] == quota


def test_an_alert_emails_only_what_is_new(customer):
    boss = customer["boss"]
    a = client.post("/api/me/alerts", headers=boss, json={"query": "الكراء", "lang": "fr",
                                                          "frequency": "daily"}).json()
    with db.session() as s:
        s.get(library.Alert, a["id"]).last_sent_at = dt.datetime.utcnow() - dt.timedelta(days=2)
        s.commit()
    _ruling("new_001")
    _ruling("old_001", created=dt.datetime.utcnow() - dt.timedelta(days=30))
    assert worker.run_alerts(dt.datetime.utcnow()) == 1
    body = _last_mail("boss@firm.ma", "alert")[1]
    assert "/fr/ruling/new_001" in body and "old_001" not in body
    assert worker.run_alerts(dt.datetime.utcnow()) == 0                 # not twice the same day


def test_renewal_reminders_go_once(customer):
    with db.session() as s:
        x = s.scalar(select(billing.Subscription).where(billing.Subscription.org_id == customer["org_id"],
                                                        billing.Subscription.status == "active"))
        x.ends_on = dt.date.today() + dt.timedelta(days=15)
        s.commit()
    assert worker.renewal_reminders(dt.date.today()) >= 1
    assert worker.renewal_reminders(dt.date.today()) == 0


def test_a_second_factor_is_asked_for_and_checked():
    h = _staff("mfa@x.com", "sales")
    secret = client.post("/api/account/mfa/setup", headers=h).json()["secret"]
    assert client.post("/api/account/mfa/enable", headers=h, json={"code": "000000"}).status_code == 400
    assert client.post("/api/account/mfa/enable", headers=h,
                       json={"code": authmod.totp(secret)}).status_code == 200
    r = client.post("/api/auth/login", json={"email": "mfa@x.com", "password": PW})
    assert r.status_code == 401 and r.json()["detail"] == {"mfa": True}
    r = client.post("/api/auth/login", json={"email": "mfa@x.com", "password": PW, "code": "123456"})
    assert r.status_code == 401
    _login("mfa@x.com", code=authmod.totp(secret))


def test_a_password_reset_ends_every_session():
    h = _staff("reset@x.com", "editor")
    assert client.post("/api/auth/forgot", json={"email": "nobody@x.com"}).json() == {"ok": True}
    client.post("/api/auth/forgot", json={"email": "reset@x.com", "language": "fr"})
    token = _token_from(_last_mail("reset@x.com", "password_reset")[1])
    assert client.post(f"/api/auth/reset/{token}", json={"password": "short"}).status_code == 400
    assert client.post(f"/api/auth/reset/{token}", json={"password": "a-new-long-password"}).status_code == 200
    assert client.get("/api/account", headers=h).status_code == 401
    _login("reset@x.com", "a-new-long-password")
    assert client.post(f"/api/auth/reset/{token}", json={"password": "a-new-long-password"}).status_code == 404


def test_support_threads_are_seen_by_the_firm_and_the_team(customer):
    clerk = _login("clerk@firm.ma")
    t = client.post("/api/support", headers=clerk, json={"subject": "Facture", "body": "Question"}).json()
    assert any(x["id"] == t["id"] for x in client.get("/api/support", headers=customer["boss"]).json())
    other = _staff("outsider@x.com", "editor")
    out = client.post(f"/api/support/{t['id']}/reply", headers=other, json={"body": "Réponse"}).json()
    assert out["status"] == "pending" and out["messages"][-1]["staff"] is True
    assert _last_mail("clerk@firm.ma", "ticket_reply")
    # another customer sees nothing of it
    db.upsert_user("lonely@y.ma", authmod.hash_password(PW), role="client")
    accounts.adopt_orphan_clients()
    assert client.get(f"/api/support/{t['id']}", headers=_login("lonely@y.ma")).status_code == 404


def test_the_admin_overview_adds_up(customer):
    d = client.get("/api/crm/overview", headers=admin()).json()
    assert d["received_month"] >= 1440000 and d["readers"] >= 1
    assert d["contracts_active"] >= 1


def test_a_request_is_converted_only_when_its_client_is_created():
    client.post("/api/public/demo", json={"full_name": "Maître Y", "email": "y@cab.ma", "language": "fr"})
    lead = next(l for l in client.get("/api/crm/leads", headers=admin()).json() if l["email"] == "y@cab.ma")
    assert lead["status"] == "new"
    out = client.post("/api/crm/clients", headers=admin(), json={
        "org": {"name": "Cabinet Y"}, "contract": {"plan_code": "individual", "amount_ht": 100000},
        "lead_id": lead["id"]}).json()
    lead = next(l for l in client.get("/api/crm/leads", headers=admin()).json() if l["email"] == "y@cab.ma")
    assert lead["status"] == "converted" and lead["org_id"] == out["org_id"]


def test_the_lists_say_what_one_client_page_says(customer):
    """The batched lists and the one-client page read access, seats and money alike."""
    rows = {o["id"]: o for o in client.get("/api/crm/orgs", headers=admin()).json()}
    for org_id, row in rows.items():
        one = client.get(f"/api/crm/orgs/{org_id}", headers=admin()).json()
        assert (row["access"], row["seats"], row["paid"], row["owed"]) == \
               (one["access"], one["seats"], one["paid"], one["owed"])
    # a client made today is new, not yet quiet or at risk
    assert rows[customer["org_id"]]["health"] in ("new", "active")


def test_owed_money_reads_the_same_on_every_page(customer):
    d = client.get("/api/crm/overview", headers=admin()).json()
    m = client.get("/api/crm/money", headers=admin()).json()
    assert (m["owed"], m["received_month"]) == (d["owed"], d["received_month"])
    orgs = client.get("/api/crm/orgs", headers=admin()).json()
    assert d["owed"] == sum(o["owed"] for o in orgs)
    # every centime owed is in an alert: awaiting a transfer, or a balance
    al = client.get("/api/crm/alerts", headers=admin()).json()
    flagged = sum(a["count"] for a in al if a["kind"] in ("awaiting_payment", "balance"))
    assert flagged == d["owed"]
