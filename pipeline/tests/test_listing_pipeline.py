"""End to end: duplicate copies, records that survive resumed runs, the listing
fields, the empty-output hold, and the two maintenance passes. No network.
"""
from __future__ import annotations

import csv
import dataclasses
import json
from pathlib import Path

from docx import Document

from anonymizer.config import Settings
from anonymizer.llm.base import BatchResult, DocIdentity, DocRef, DocumentAI, PIIEntity
from anonymizer.pipeline import Pipeline, RunOptions
from test_pipeline_records import FILLER

HEADER = "قرار محكمة النقض عدد 53 الصادر بتاريخ 08 فبراير 2022 في الملف الشرعي رقم 2020/2/2/516"


class FakeAI(DocumentAI):
    """Flags one name; classifies by topic, the way the real model does."""
    category = "اجتماعية"

    def process_pdf(self, pdf_bytes: bytes, page_count: int = 1) -> BatchResult:
        return self._r()

    def process_text(self, text: str) -> BatchResult:
        return self._r()

    def _r(self) -> BatchResult:
        return BatchResult(
            pii=[PIIEntity("محمد العلوي", "name")],
            court="محكمة النقض", category=self.category,
            identity=DocIdentity(level="نقض", court="محكمة النقض", decision_no="53/1",
                                 date="08 فبراير 2022", file_no="2020/2/2/516"),
            refs=[DocRef(court="محكمة الاستئناف بطنجة", decision_no="12", file_no="44/1601/2019")],
            input_tokens=100, output_tokens=50,
        )


def _docx(path: Path, *lines: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    d = Document()
    for line in lines:
        d.add_paragraph(line)
    d.save(str(path))


def _ruling(path: Path, extra: str = "") -> None:
    _docx(path, HEADER, "حضر محمد العلوي " + extra, FILLER)


def _settings(tmp: Path, folder: str = "in", out: str = "out", db: str = "cases.db") -> Settings:
    return dataclasses.replace(
        Settings.load(), provider="fake", api_key="x", soffice_path="",
        input_dir=tmp / folder, output_dir=tmp / out,
        state_file=tmp / out / ".state.json", db_path=tmp / db,
        max_workers=1, budget_usd=0, reader_processes=0, corpus="", default_category="",
    )


def _csv(path: Path) -> list[dict]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def test_byte_identical_copies_are_processed_once(tmp_path):
    s = _settings(tmp_path)
    _ruling(s.input_dir / "a.docx")
    (s.input_dir / "a_1.docx").write_bytes((s.input_dir / "a.docx").read_bytes())
    _ruling(s.input_dir / "b.docx", "وشخص آخر")

    calls = []

    class Counting(FakeAI):
        def process_text(self, text):
            calls.append(text)
            return super().process_text(text)

    Pipeline(s, Counting()).run(RunOptions())
    assert len(calls) == 2, "the copy must not be sent to the model"
    assert _csv(s.output_dir / "duplicates.csv") == [{"file": "a_1", "same_as": "a"}]
    records = json.loads((s.output_dir / "records.json").read_text(encoding="utf-8"))
    assert sorted(r["doc_id"] for r in records) == ["a", "b"]


def test_a_copy_already_processed_by_an_earlier_run_is_retired(tmp_path):
    s = _settings(tmp_path)
    _ruling(s.input_dir / "a.docx")
    (s.input_dir / "a_1.docx").write_bytes((s.input_dir / "a.docx").read_bytes())
    p = Pipeline(s, FakeAI())
    # Simulate the old behaviour: both copies delivered before copies were detected.
    p._drop_duplicates = lambda docs: (setattr(p, "_hashes", {}) or docs)
    p.run(RunOptions())
    assert (s.output_dir / "a_1.docx").exists()

    Pipeline(s, FakeAI()).run(RunOptions())
    assert not (s.output_dir / "a_1.docx").exists()
    assert (s.output_dir / "_duplicates" / "a_1.docx").exists(), "moved aside, not deleted"
    assert [r["file"] for r in _csv(s.output_dir / "listing.csv")] == ["a"]


def test_a_copy_of_a_ruling_from_another_folder_sharing_the_database_is_skipped(tmp_path):
    first = _settings(tmp_path, folder="f1", out="o1")
    _ruling(first.input_dir / "x.docx")
    Pipeline(first, FakeAI()).run(RunOptions())

    second = _settings(tmp_path, folder="f2", out="o2")
    second.input_dir.mkdir(parents=True)
    (second.input_dir / "renamed.docx").write_bytes((first.input_dir / "x.docx").read_bytes())
    Pipeline(second, FakeAI()).run(RunOptions())
    assert _csv(second.output_dir / "duplicates.csv") == [{"file": "renamed", "same_as": "x"}]
    assert _csv(second.output_dir / "listing.csv") == []
    # the first folder's exports are untouched by the second folder's run
    assert [r["file"] for r in _csv(first.output_dir / "listing.csv")] == ["x"]


def test_records_cover_every_document_across_resumed_runs(tmp_path):
    """records.json used to describe only the last run: 11 of 3,239 rulings."""
    s = _settings(tmp_path)
    _ruling(s.input_dir / "a.docx")
    _ruling(s.input_dir / "b.docx", "وآخر")
    Pipeline(s, FakeAI()).run(RunOptions(limit=1))
    Pipeline(s, FakeAI()).run(RunOptions())
    records = json.loads((s.output_dir / "records.json").read_text(encoding="utf-8"))
    assert sorted(r["doc_id"] for r in records) == ["a", "b"]
    # Not just listed: the first run's record keeps everything it knew, rather than
    # the thin record that can be rebuilt from the state file alone.
    first = next(r for r in records if r["doc_id"] == "a")
    assert first["read_method"] == "text"
    assert first["fields"]["رقم القرار"]["value"] == "53"


def test_listing_fields_come_out_in_portal_form(tmp_path):
    s = _settings(tmp_path)
    _ruling(s.input_dir / "a.docx")
    Pipeline(s, FakeAI()).run(RunOptions())
    [row] = _csv(s.output_dir / "listing.csv")
    assert row["رقم القرار"] == "53"               # "53/1" -> decision without section
    assert row["تاريخ القرار"] == "08/02/2022"
    assert row["السنة"] == "2022"
    assert row["المدينة"] == "طنجة"                # the lower court, not Rabat
    # The model said "social" (the topic); the ruling says it is a personal-status file.
    assert row["الغرفة"] == "أحوال شخصية"
    rec = json.loads((s.output_dir / "records.json").read_text(encoding="utf-8"))[0]
    assert rec["category"] == {"value": "أحوال شخصية", "source": "ruling-text"}
    assert rec["city"] == "طنجة" and rec["year"] == "2022"


def test_an_almost_empty_output_is_held_not_delivered(tmp_path):
    s = _settings(tmp_path)
    _docx(s.input_dir / "blank.docx", "حضر محمد العلوي")    # a title's worth of text
    Pipeline(s, FakeAI()).run(RunOptions())
    assert not (s.output_dir / "blank.docx").exists()
    assert (s.output_dir / "_quarantine" / "blank.docx").exists()
    [q] = _csv(s.output_dir / "quarantine.csv")
    assert q["reason"] == "empty"


def test_quarantine_csv_never_quotes_the_surviving_names(tmp_path):
    s = _settings(tmp_path)

    class Leaky(FakeAI):
        def _r(self):
            r = super()._r()
            r.pii = [PIIEntity("زيد", "name")]   # flagged, but the text also has a variant
            return r

    _docx(s.input_dir / "leaky.docx", HEADER, "حضر زيد", FILLER)
    p = Pipeline(s, Leaky())
    import anonymizer.pipeline as pl
    real = pl.redact_text
    pl.redact_text = lambda text, ents, tok: (text, real(text, [], tok)[1])  # redactor misses it
    try:
        p.run(RunOptions())
    finally:
        pl.redact_text = real
    [q] = _csv(s.output_dir / "quarantine.csv")
    assert q["reason"] == "leak"
    assert "زيد" not in q["detail"]
    sidecar = (s.output_dir / "_records" / "leaky.json").read_text(encoding="utf-8")
    assert "زيد" not in sidecar


def test_refresh_corrects_the_chamber_and_its_title_without_the_model(tmp_path):
    s = _settings(tmp_path)
    _ruling(s.input_dir / "a.docx")

    class OldRules(FakeAI):
        pass

    p = Pipeline(s, OldRules())
    import anonymizer.pipeline as pl
    real = pl.chamber_from_text
    pl.chamber_from_text = lambda text: ""          # the pipeline before this change
    try:
        p.run(RunOptions())
    finally:
        pl.chamber_from_text = real
    out = s.output_dir / "a.docx"
    assert Document(str(out)).paragraphs[0].text == "محكمة النقض — اجتماعية"

    class NoCalls(FakeAI):
        def process_text(self, text):
            raise AssertionError("refresh must not call the model")

    q = Pipeline(s, NoCalls())
    counts = q.refresh_listing()
    q.export()
    assert counts["chamber_changed"] == 1 and counts["retitled"] == 1
    assert Document(str(out)).paragraphs[0].text == "محكمة النقض — أحوال شخصية"
    [row] = _csv(s.output_dir / "listing.csv")
    assert row["الغرفة"] == "أحوال شخصية"


def test_requeue_sets_aside_files_with_words_cut_open(tmp_path):
    s = _settings(tmp_path)
    _ruling(s.input_dir / "a.docx")
    Pipeline(s, FakeAI()).run(RunOptions())
    out = s.output_dir / "a.docx"
    d = Document(str(out))
    d.add_paragraph("قضت XXXXXXXحكمة بما يلي")          # the old over-redaction
    d.save(str(out))
    (s.output_dir / "_records" / "a.json").unlink()      # i.e. written by the old code

    counts = Pipeline(s, FakeAI()).requeue_damaged()
    assert counts["cut-words"] == 1
    assert not out.exists() and (s.output_dir / "_superseded" / "a.docx").exists()
    Pipeline(s, FakeAI()).run(RunOptions())
    assert out.exists(), "the next normal run reprocesses it"


def test_a_partial_transcription_is_held_not_delivered(tmp_path):
    """Six pages in, only the ruling's closing lines out: 5 such files were delivered."""
    s = _settings(tmp_path)
    _ruling(s.input_dir / "scan.docx")

    class PartialOCR(FakeAI):
        pass

    p = Pipeline(s, PartialOCR())
    import anonymizer.pipeline as pl
    real = pl.iter_work_batches

    class _Batch:
        kind, pdf_bytes, page_count = "image", b"", 6

    p.provider.process_pdf = lambda b, n: BatchResult(
        # ~500 characters for six pages: above the empty line, far below a ruling.
        pages=[HEADER + " لهذه الأسباب قضت محكمة النقض برفض الطلب. " + FILLER[:330]] + [""] * 5,
        pii=[], court="محكمة النقض", category="أحوال شخصية")
    pl.iter_work_batches = lambda *a, **k: iter([_Batch()])
    try:
        p.run(RunOptions())
    finally:
        pl.iter_work_batches = real
    assert not (s.output_dir / "scan.docx").exists()
    [q] = _csv(s.output_dir / "quarantine.csv")
    assert q["reason"] == "partial"


def test_the_listing_shows_only_delivered_rulings(tmp_path):
    """A held file has no document behind it for a reader to open."""
    s = _settings(tmp_path)
    _ruling(s.input_dir / "good.docx")
    _docx(s.input_dir / "blank.docx", "حضر محمد العلوي")      # held as empty
    Pipeline(s, FakeAI()).run(RunOptions())
    assert [r["file"] for r in _csv(s.output_dir / "listing.csv")] == ["good"]


def test_refresh_rebuilds_a_ruling_missing_from_the_database(tmp_path):
    s = _settings(tmp_path)
    _ruling(s.input_dir / "a.docx")
    Pipeline(s, FakeAI()).run(RunOptions())
    p = Pipeline(s, FakeAI())
    p.casedb.remove("a")                     # as if processed before the DB moved here
    counts = p.refresh_listing()
    p.export()
    assert counts.get("rows_rebuilt") == 1
    [row] = _csv(s.output_dir / "listing.csv")
    assert row["file"] == "a" and row["الغرفة"] == "أحوال شخصية" and row["رقم القرار"] == "53"


def test_requeue_leaves_current_output_alone(tmp_path):
    """A full name glued to its neighbour in today's output is the intended redaction."""
    s = _settings(tmp_path)
    _ruling(s.input_dir / "a.docx")
    Pipeline(s, FakeAI()).run(RunOptions())
    out = s.output_dir / "a.docx"
    d = Document(str(out))
    d.add_paragraph("حضر XXXXXXXبمقال افتتاحي")
    d.save(str(out))
    assert Pipeline(s, FakeAI()).requeue_damaged()["cut-words"] == 0
    assert out.exists()


def test_sample_draws_a_repeatable_random_subset(tmp_path):
    s = _settings(tmp_path)
    for i in range(12):
        _ruling(s.input_dir / f"r{i:02d}.docx", f"رقم {i}")
    seen = []

    class Recording(FakeAI):
        def process_text(self, text):
            seen.append(text)
            return super().process_text(text)

    Pipeline(s, Recording()).run(RunOptions(sample=4))
    done = sorted(d for d, st in json.loads(
        (s.output_dir / ".state.json").read_text(encoding="utf-8")).items() if st["status"] == "done")
    assert len(done) == 4
    assert done != ["r00", "r01", "r02", "r03"], "not simply the first four by name"


def test_state_survives_windows_briefly_locking_the_file(tmp_path, monkeypatch):
    """A scanner holding the state file made one replace fail and killed a run."""
    import os
    from anonymizer.core import state as state_mod

    real, calls = os.replace, []

    def flaky(src, dst):
        calls.append(1)
        if len(calls) <= 2:
            raise PermissionError(5, "Access is denied")
        return real(src, dst)

    monkeypatch.setattr(state_mod.os, "replace", flaky)
    st = state_mod.State(tmp_path / ".state.json")
    st.mark_duplicates({f"copy{i}": "orig" for i in range(1000)})
    assert len(calls) == 3, "one save for all copies, retried past the lock"
    assert json.loads((tmp_path / ".state.json").read_text(encoding="utf-8"))["copy999"]["status"] == "duplicate"


def test_a_legacy_record_carries_the_listing_fields_into_records_json(tmp_path):
    """records.json is what the platform imports; rebuilt records had no fields."""
    s = _settings(tmp_path)
    _ruling(s.input_dir / "a.docx")
    Pipeline(s, FakeAI()).run(RunOptions())
    (s.output_dir / "_records" / "a.json").unlink()      # processed before sidecars
    p = Pipeline(s, FakeAI())
    p.export()
    [rec] = json.loads((s.output_dir / "records.json").read_text(encoding="utf-8"))
    f = rec["fields"]
    assert f["رقم القرار"]["value"] == "53" and f["تاريخ القرار"]["value"] == "08/02/2022"
    assert f["المدينة"]["value"] == "طنجة" and f["الغرفة"]["value"] == "أحوال شخصية"


def test_a_name_flagged_in_a_later_batch_is_removed_from_earlier_pages(tmp_path):
    """Long rulings are read in batches; each page used to get only its batch's names.

    Here the name is flagged only while reading the second batch, but it also
    stands on the first page.
    """
    s = _settings(tmp_path)
    _ruling(s.input_dir / "long.docx")

    class _Batch:
        def __init__(self, text):
            self.kind, self.pages, self.page_count = "text", [text], 1

    # Two full pages each (short pages would be held as a partial transcription).
    first = HEADER + " حضر خالد البناني أمام المحكمة. " + FILLER * 2
    second = "وبعد المداولة أكد خالد البناني طلبه. " + FILLER * 2

    class SecondBatchOnly(FakeAI):
        def process_text(self, text):
            r = super().process_text(text)
            r.pii = [PIIEntity("خالد البناني", "name")] if text.startswith("وبعد") else []
            return r

    import anonymizer.pipeline as pl
    real = pl.iter_work_batches
    pl.iter_work_batches = lambda *a, **k: iter([_Batch(first), _Batch(second)])
    try:
        Pipeline(s, SecondBatchOnly()).run(RunOptions())
    finally:
        pl.iter_work_batches = real
    out = s.output_dir / "long.docx"
    held = _csv(s.output_dir / "quarantine.csv")
    assert out.exists(), f"not held: the name is removed everywhere ({held})"
    assert "البناني" not in "\n".join(p.text for p in Document(str(out)).paragraphs)


def test_reading_in_separate_processes_gives_byte_identical_output(tmp_path):
    """The reader processes are a speed-up only: nothing delivered may change."""
    import dataclasses
    import hashlib

    outputs = {}
    for procs in (0, 2):
        s = dataclasses.replace(_settings(tmp_path / f"p{procs}"), reader_processes=procs)
        for n in ("a", "b", "c"):
            _ruling(s.input_dir / f"{n}.docx", n)
        Pipeline(s, FakeAI()).run(RunOptions())
        outputs[procs] = {f.name: hashlib.sha256(f.read_bytes()).hexdigest()
                          for f in sorted(s.output_dir.glob("*.docx"))}
    assert len(outputs[0]) == 3
    assert outputs[0] == outputs[2]


def test_state_writes_are_throttled_but_nothing_is_lost(tmp_path):
    from anonymizer.core.state import DocRecord, State

    st = State(tmp_path / "s.json", save_interval=3600)
    for i in range(50):
        st.update(f"d{i}", DocRecord(status="done"))
    on_disk = json.loads((tmp_path / "s.json").read_text(encoding="utf-8"))
    assert len(on_disk) == 1, "only the first update is written inside the interval"
    st.flush()
    assert len(json.loads((tmp_path / "s.json").read_text(encoding="utf-8"))) == 50
